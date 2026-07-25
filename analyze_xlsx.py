import openpyxl
import warnings
warnings.filterwarnings('ignore')

fp = 'c:\\Users\\Roman\\YandexDisk\\ProjectSQL\\Broker_Report\\reports\\report963464ae-b003-a5d1-c98b-2f4abc8d75ed.xlsx'
wb = openpyxl.load_workbook(fp, data_only=True)

print('Sheets:', wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'Sheet: {sn}, rows={ws.max_row}, cols={ws.max_column}')
    for r in range(1, min(ws.max_row+1, 200)):
        vals = []
        for c in range(1, min(ws.max_column+1, 12)):
            v = ws.cell(r, c).value
            if v is not None:
                s = str(v).replace('\n', ' ')[:50]
                vals.append(f'C{c}={s}')
        if vals:
            sep = ' | '
            print(f'  R{r}: {sep.join(vals)}')
